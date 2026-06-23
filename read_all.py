import openpyxl

def print_sheet_all(path, sheet_name, label):
    print(f"\n{'='*120}")
    print(f"=== {label}")
    print(f"=== File: {path}")
    print(f"{'='*120}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    print(f"Sheet: {sheet_name} | Rows: {ws.max_row} | Cols: {ws.max_column}")
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=False), start=1):
        vals = []
        for cell in row:
            vals.append(str(cell.value) if cell.value is not None else "")
        print(f"Row {row_idx:>3}: " + " | ".join(vals))
    wb.close()

# 1. MONOPUNTO - RECTYAFIL
print_sheet_all(
    r"C:\Users\PC41\Desktop\Dev\Actualización precios\MONOPUNTO\212 Datos Monopunto RECTYAFIL.xlsx",
    "Plantilla",
    "1. MONOPUNTO: RECTYAFIL - ALL ROWS"
)

# 2. MULTIPUNTO - ABULA GOLF
print_sheet_all(
    r"C:\Users\PC41\Desktop\Dev\Actualización precios\MULTIPUNTO\Datos Multipunto ABULA GOLF.xlsx",
    "Plantilla",
    "2. MULTIPUNTO: ABULA GOLF - ALL ROWS"
)

# 3. NEXUS oferta
nexus_path = r"C:\Users\PC41\Desktop\Dev\Actualización precios\Ofertas\Nexus - PETICION ELECTRICIDAD JUNIO.xlsx"
print(f"\n{'='*120}")
print("=== 3. NEXUS Oferta File")
print(f"{'='*120}")
wb = openpyxl.load_workbook(nexus_path, data_only=True)
print(f"Sheets: {wb.sheetnames}")
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} | Rows: {ws.max_row} | Cols: {ws.max_column}")

    # Print all column headers (first row)
    print("\n--- ALL COLUMN HEADERS (Row 1):")
    print(f"Total columns: {ws.max_column}")
    for col_idx in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col_idx).value
        print(f"  Col {col_idx:>2}: {v}")

    # Print first 5 rows
    print(f"\n--- First 5 rows:")
    for row_idx in range(1, min(6, ws.max_row + 1)):
        vals = []
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            vals.append(str(v) if v is not None else "")
        print(f"Row {row_idx:>3}: " + " | ".join(vals))

    # Find RECTYAFIL rows
    print(f"\n--- Searching for 'RECTYAFIL' rows:")
    rectyafil_found = False
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None and "RECTYAFIL" in str(v).upper():
                rectyafil_found = True
                vals = []
                for c in range(1, ws.max_column + 1):
                    cv = ws.cell(row=row_idx, column=c).value
                    vals.append(str(cv) if cv is not None else "")
                print(f"Row {row_idx:>3}: " + " | ".join(vals))
                break
    if not rectyafil_found:
        print("  (Not found in this sheet)")

    # Find ABULA GOLF rows
    print(f"\n--- Searching for 'ABULA GOLF' rows:")
    abula_found = False
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None and "ABULA" in str(v).upper() and "GOLF" in str(v).upper():
                abula_found = True
                vals = []
                for c in range(1, ws.max_column + 1):
                    cv = ws.cell(row=row_idx, column=c).value
                    vals.append(str(cv) if cv is not None else "")
                print(f"Row {row_idx:>3}: " + " | ".join(vals))
                break
    if not abula_found:
        print("  (Not found in this sheet)")
wb.close()

# 4. ELEIA oferta - ABULA GOLF
eleia_path = r"C:\Users\PC41\Desktop\Dev\Actualización precios\Ofertas\ELEIA - PETICION ELECTRICIDAD JUNIO PRECIOS.xlsx"
print(f"\n{'='*120}")
print("=== 4. ELEIA Oferta File - ABULA GOLF rows")
print(f"{'='*120}")
wb = openpyxl.load_workbook(eleia_path, data_only=True)
print(f"Sheets: {wb.sheetnames}")
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} | Rows: {ws.max_row} | Cols: {ws.max_column}")

    # Print all column headers
    print(f"\n--- ALL COLUMN HEADERS (Row 1):")
    for col_idx in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col_idx).value
        print(f"  Col {col_idx:>2}: {v}")

    # Find ABULA GOLF rows
    print(f"\n--- Searching for 'ABULA GOLF' rows:")
    abula_found = False
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None and "ABULA" in str(v).upper() and "GOLF" in str(v).upper():
                abula_found = True
                vals = []
                for c in range(1, ws.max_column + 1):
                    cv = ws.cell(row=row_idx, column=c).value
                    vals.append(str(cv) if cv is not None else "")
                print(f"Row {row_idx:>3}: " + " | ".join(vals))
                break
    if not abula_found:
        # Try just ABULA
        print("  Trying 'ABULA' only...")
        for row_idx in range(1, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                v = ws.cell(row=row_idx, column=col_idx).value
                if v is not None and "ABULA" in str(v).upper():
                    abula_found = True
                    vals = []
                    for c in range(1, ws.max_column + 1):
                        cv = ws.cell(row=row_idx, column=c).value
                        vals.append(str(cv) if cv is not None else "")
                    print(f"Row {row_idx:>3}: " + " | ".join(vals))
                    break
        if not abula_found:
            print("  (Not found)")
wb.close()

print("\n\nDONE.")
