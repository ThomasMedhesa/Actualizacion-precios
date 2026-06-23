import openpyxl
from collections import defaultdict


def fmt(v):
    if v is None:
        return "N/A"
    try:
        return f"{float(v):.4f}"
    except:
        return str(v)[:10]


def fmt2(v):
    if v is None:
        return "N/A"
    try:
        return f"{float(v):.3f}"
    except:
        return str(v)[:10]


def analyze_file(path, label):
    print(f"\n{'='*120}")
    print(f"=== {label}")
    print(f"{'='*120}")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["PETICION OFERTAS"]

    blocks = []

    for r in range(8, ws.max_row + 1):
        cliente = ws.cell(row=r, column=2).value
        codigo = ws.cell(row=r, column=4).value
        cups = ws.cell(row=r, column=9).value
        tarifa = ws.cell(row=r, column=10).value
        p1 = ws.cell(row=r, column=26).value
        p2 = ws.cell(row=r, column=27).value
        p3 = ws.cell(row=r, column=28).value
        p4 = ws.cell(row=r, column=29).value
        p5 = ws.cell(row=r, column=30).value
        p6 = ws.cell(row=r, column=31).value
        fee = ws.cell(row=r, column=32).value
        desvios = ws.cell(row=r, column=33).value

        if cliente is not None and str(cliente).strip():
            current_client = str(cliente).strip()
            current_codigo = str(codigo).strip() if codigo is not None else ""
            blocks.append({
                "row": r,
                "cliente": current_client,
                "codigo": current_codigo,
                "cups": str(cups).strip() if cups is not None else "",
                "tarifa": str(tarifa).strip() if tarifa is not None else "",
                "p1": p1,
                "p2": p2,
                "p3": p3,
                "p4": p4,
                "p5": p5,
                "p6": p6,
                "fee": fee,
                "desvios": desvios,
                "period_count": 0,
            })
        else:
            if blocks:
                blocks[-1]["period_count"] += 1
                b = blocks[-1]
                if p1 is not None and b["p1"] is None:
                    b["p1"] = p1
                if p2 is not None and b["p2"] is None:
                    b["p2"] = p2
                if p3 is not None and b["p3"] is None:
                    b["p3"] = p3
                if p4 is not None and b["p4"] is None:
                    b["p4"] = p4
                if p5 is not None and b["p5"] is None:
                    b["p5"] = p5
                if p6 is not None and b["p6"] is None:
                    b["p6"] = p6
                if fee is not None and b["fee"] is None:
                    b["fee"] = fee
                if desvios is not None and b["desvios"] is None:
                    b["desvios"] = desvios

    print(f"Total data rows: {ws.max_row - 7}")
    print(f"Supply point blocks: {len(blocks)}")
    print(f"Unique Códigos: {len(set(b['codigo'] for b in blocks))}")
    print()

    # --- Group by CLIENTE ---
    cliente_to_blocks = defaultdict(list)
    for b in blocks:
        cliente_to_blocks[b["cliente"]].append(b)

    monopunto_clients = {k: v for k, v in cliente_to_blocks.items() if len(v) == 1}
    multipunto_clients = {k: v for k, v in cliente_to_blocks.items() if len(v) > 1}

    # --- Group by CÓDIGO ---
    codigo_to_blocks = defaultdict(list)
    for b in blocks:
        codigo_to_blocks[b["codigo"]].append(b)

    monopunto_cod = {k: v for k, v in codigo_to_blocks.items() if len(v) == 1}
    multipunto_cod = {k: v for k, v in codigo_to_blocks.items() if len(v) > 1}

    # Clean: remove empty-codigo entries
    actual_codigos = {k: v for k, v in codigo_to_blocks.items() if k != ""}
    monopunto_actual = {k: v for k, v in actual_codigos.items() if len(v) == 1}
    print(f"--- Grouping by CÓDIGO (excluding empty Código): ---")
    print(f"Códigos with count=1 (unique supply point): {len(monopunto_actual)}")
    print(f"Códigos with count>1: {len({k for k in actual_codigos if len(actual_codigos[k])>1})}")

    # Show the count>1 Códigos
    for codigo in sorted(actual_codigos.keys()):
        if len(actual_codigos[codigo]) > 1:
            print(f"  Código '{codigo}' appears {len(actual_codigos[codigo])}x -> clients: {', '.join(sorted(set(b['cliente'] for b in actual_codigos[codigo])))}")

    print()
    print(f"--- Grouping by CLIENTE name: ---")
    print(f"Clients with 1 supply point (MONOPUNTO): {len(monopunto_clients)}")
    print(f"Clients with >1 supply points (MULTIPUNTO): {len(multipunto_clients)}")
    print()

    print("=== MONOPUNTO CLIENTS (1 Código) ===")
    for cli in sorted(monopunto_clients.keys()):
        b = monopunto_clients[cli][0]
        print(f"  {cli:50s} Código={b['codigo']:12s} CUPS={b['cups']:30s} Tarifa={b['tarifa']:10s} Periodos={b['period_count']:2d}  P1={fmt(b['p1']):10s} P2={fmt(b['p2']):10s} P3={fmt(b['p3']):10s} P4={fmt(b['p4']):10s} P5={fmt(b['p5']):10s} P6={fmt(b['p6']):10s} FEE={fmt2(b['fee']):8s} Desv={fmt2(b['desvios']):8s}")

    print()
    print("=== MULTIPUNTO CLIENTS (>1 Códigos) ===")
    for cli in sorted(multipunto_clients.keys()):
        bloques = multipunto_clients[cli]
        codigos = sorted(set(b["codigo"] for b in bloques))
        print(f"  {cli:50s} {len(codigos)} Códigos: {', '.join(codigos)}")
        for b in bloques:
            print(f"    -> Código={b['codigo']:12s} CUPS={b['cups']:30s} Tarifa={b['tarifa']:10s} P1={fmt(b['p1']):10s} FEE={fmt2(b['fee']):8s}")

    print()
    # Cross-reference with known template clients
    known_mono = ["RECTYAFIL", "CECOSA", "SANTA RITA HARINAS"]
    known_multi = ["ABULA GOLF", "AGROPECUARIA NAVALOSHACES", "ALCALA 534"]
    print("=== Cross-reference with MONOPUNTO/MULTIPUNTO template files ===")
    for km in known_mono:
        for b in blocks:
            if km.upper() in b["cliente"].upper():
                print(f"  [MONOPUNTO template] {km:30s} -> {b['cliente'][:40]:42s} Código={b['codigo']:12s} CUPS={b['cups']:30s} Tarifa={b['tarifa']}")
                break
    for km in known_multi:
        for b in blocks:
            if km.upper() in b["cliente"].upper():
                print(f"  [MULTIPUNTO template] {km:30s} -> {b['cliente'][:40]:42s} Código={b['codigo']:12s} CUPS={b['cups']:30s} Tarifa={b['tarifa']}")

    wb.close()
    return blocks


eleia_path = r"C:\Users\PC41\Desktop\Dev\Actualización precios\Ofertas\ELEIA - PETICION ELECTRICIDAD JUNIO PRECIOS.xlsx"
iber_path = r"C:\Users\PC41\Desktop\Dev\Actualización precios\Ofertas\Iberdrola - PETICION ELECTRICIDAD JUNIO1.xlsx"

analyze_file(eleia_path, "1. ELEIA - PETICION ELECTRICIDAD JUNIO PRECIOS.xlsx")
analyze_file(iber_path, "2. Iberdrola - PETICION ELECTRICIDAD JUNIO1.xlsx")

print("\nDONE.")
